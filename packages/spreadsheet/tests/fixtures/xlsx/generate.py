"""Regenerate small, independent OOXML fixtures with XlsxWriter and openpyxl.

Only fixture generation needs these Python packages; normal tests read the checked-in files.
All names, numbers, images and comments below are synthetic.
"""
from pathlib import Path
from datetime import datetime
import io
import struct
import zlib
import xlsxwriter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).parent
def png_chunk(kind, data):
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

pixels = b"".join(b"\0" + b"".join(bytes((33, 110, 57) if x < 80 else (242, 177, 52)) for x in range(160)) for _ in range(80))
PNG = b"\x89PNG\r\n\x1a\n" + png_chunk(b"IHDR", struct.pack(">IIBBBBB", 160, 80, 8, 2, 0, 0, 0)) + png_chunk(b"IDAT", zlib.compress(pixels)) + png_chunk(b"IEND", b"")

def xlsxwriter_fixture(path, warnings=False):
    book = xlsxwriter.Workbook(path)
    book.set_properties({"title": "LikeX import fixture", "author": "LikeX tests", "created": datetime(2026, 1, 1)})
    sheet = book.add_worksheet("売上")
    title = book.add_format({"bold": True, "font_size": 18, "font_color": "#216E39", "bg_color": "#EAF3EE", "align": "center"})
    date = book.add_format({"num_format": "yyyy-mm-dd"})
    text = book.add_format({"num_format": "@"})
    sheet.merge_range("A1:D1", "2026年度 売上サンプル", title)
    sheet.set_row(0, 34)
    sheet.set_column("A:A", 24)
    sheet.set_column("B:D", 15)
    sheet.add_table("A3:D6", {"name": "Sales", "columns": [{"header": x} for x in ["商品", "数量", "単価", "小計"]], "data": [["りんご", 2, 120, 240], ["みかん", 3, 80, 240], ["ぶどう", 1, 500, 500]]})
    sheet.write("C8", "合計")
    sheet.write_formula("D8", "=SUM(D4:D6)", None, 980)
    sheet.write_datetime("A10", datetime(2026, 9, 16), date)
    sheet.data_validation("A10", {"validate": "date", "criteria": "between", "minimum": datetime(2026, 1, 1), "maximum": datetime(2026, 12, 31)})
    sheet.write_string("B10", "00123", text)
    sheet.write_string("C10", "=literal", text)
    sheet.write_boolean("D10", True)
    sheet.data_validation("B12:B14", {"validate": "list", "source": ["未着手", "進行中", "完了"]})
    sheet.data_validation("C12:C14", {"validate": "integer", "criteria": "between", "minimum": 1, "maximum": 10})
    sheet.write_comment("A4", "テスト用コメント", {"author": "Sample"})
    book.define_name("売上明細", "='売上'!$A$4:$D$6")
    sheet.insert_image("F2", "sample.png", {"image_data": io.BytesIO(PNG), "x_scale": 0.5, "y_scale": 0.5, "description": "テスト画像"})
    sheet.insert_textbox("F6", "文字入りテキストボックス\n取り込みサンプル", {"width": 220, "height": 80, "fill": {"color": "#EAF3EE"}, "line": {"color": "#216E39"}, "font": {"color": "#216E39"}})
    empty = book.add_worksheet("空シート")
    empty.write("A1", "2枚目も保持")
    if warnings:
        sheet.write_formula("A16", "=UNSUPPORTED_FUNCTION(1)", None, 42)
        chart = book.add_chart({"type": "column"})
        chart.add_series({"values": "='売上'!$D$4:$D$6"})
        sheet.insert_chart("F12", chart)
    book.close()

xlsxwriter_fixture(ROOT / "xlsxwriter-sample.xlsx")
xlsxwriter_fixture(ROOT / "xlsxwriter-warnings.xlsx", warnings=True)
book = Workbook()
book.properties.creator = "LikeX tests"
book.properties.created = book.properties.modified = datetime(2026, 1, 1)
sheet = book.active
sheet.title = "Inline strings"
sheet["A1"] = "00123"
sheet["A2"] = "=SUM(B1:B2)"
sheet["B1"] = 2
sheet["B2"] = 3
sheet["C1"] = datetime(2026, 9, 16)
sheet["C1"].number_format = "yyyy-mm-dd"
sheet["A1"].font = Font(bold=True, color="216E39")
sheet["A1"].fill = PatternFill("solid", fgColor="EAF3EE")
sheet.merge_cells("A4:C4")
sheet["A4"] = "結合セル"
book.save(ROOT / "openpyxl-sample.xlsx")
