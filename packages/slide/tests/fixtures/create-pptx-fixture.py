"""Optional fixture regeneration: python-pptx 1.0.2 + Pillow. Runtime tests need only the checked-in PPTX."""
from pathlib import Path
from io import BytesIO
from datetime import datetime, timezone
from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image

p = Presentation()
p.slide_width, p.slide_height = Inches(13.333333), Inches(7.5)
p.core_properties.title = "PowerPoint compatibility fixture"
p.core_properties.author = "LikeX"
p.core_properties.last_modified_by = "LikeX"
p.core_properties.comments = "LikeX PPTX interoperability fixture generated with python-pptx"
p.core_properties.created = datetime(2026, 9, 16, tzinfo=timezone.utc)
p.core_properties.modified = datetime(2026, 9, 16, tzinfo=timezone.utc)
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = "Inherited title geometry"
s.placeholders[1].text = "Master font and theme color"
s.notes_slide.notes_text_frame.text = "Presenter notes\nSecond line"
s = p.slides.add_slide(p.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor(0xE7, 0xF0, 0xFA)
box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
p0 = box.text_frame.paragraphs[0]
p0.alignment = PP_ALIGN.RIGHT
run = p0.add_run()
run.text = "日本語 & <PowerPoint>"
run.font.name, run.font.size, run.font.bold = "Arial", Pt(24), True
run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(3), Inches(3), Inches(1.2))
shape.rotation = 15
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0x33, 0x99, 0x66)
shape.line.color.rgb = RGBColor(0x22, 0x22, 0x22)
shape.text = "Shape text"
buf = BytesIO()
Image.new("RGB", (40, 20), (255, 100, 20)).save(buf, format="PNG")
buf.seek(0)
image = s.shapes.add_picture(buf, Inches(7), Inches(2), width=Inches(3))
image.rotation = 20
package = BytesIO()
p.save(package)
package.seek(0)
with ZipFile(package) as source, ZipFile(Path(__file__).with_name("powerpoint-basic.pptx"), "w") as output:
    for entry in source.infolist():
        entry.date_time = (2026, 9, 16, 0, 0, 0)
        output.writestr(entry, source.read(entry.filename))
