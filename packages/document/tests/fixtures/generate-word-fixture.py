"""Development-only compatibility fixture, generated independently with python-docx (MIT)."""
from base64 import b64decode
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor

document = Document()
document.core_properties.title = "Independent Word fixture"
section = document.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.left_margin, section.right_margin = Mm(18), Mm(18)
document.add_heading("Project notes", level=1)
paragraph = document.add_paragraph()
run = paragraph.add_run("Styled Japanese 日本語")
run.bold, run.italic = True, True
run.font.name, run.font.size, run.font.color.rgb = "Arial", Pt(16), RGBColor.from_string("123456")
document.add_paragraph("First item", style="List Bullet")
document.add_paragraph("Second item", style="List Bullet")
table = document.add_table(rows=3, cols=2)
table.style = "Table Grid"
for row, values in zip(table.rows, [("Name", "Value"), ("Alpha", "100"), ("Beta", "200")]):
    for cell, value in zip(row.cells, values):
        cell.text = value
png = b64decode("iVBORw0KGgoAAAANSUhEUgAAACgAAAAUCAIAAABwJOjsAAAAOElEQVR4nO3NQQEAIAwDsVINSMS/BfhuBm4PGgNZ92xN8MiqxCCTWZUYY67qEmPMVV1ijLlKn8cPnj8Bn7y225YAAAAASUVORK5CYII=")
image = document.add_picture(BytesIO(png), width=Mm(60))
image._inline.docPr.set("descr", "Independent image")
document.add_page_break()
document.add_paragraph("After page break")
document.save(Path(__file__).with_name("word-basic.docx"))
